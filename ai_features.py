"""
ai_features.py — All AI/ML features for ConverterPro
ATS scoring, CV enhancement, cover letters, QR codes, email, resume parsing
"""
import os, io, re, json, base64, tempfile

try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False

try:
    import qrcode
    QR_AVAILABLE = True
except ImportError:
    QR_AVAILABLE = False

try:
    import sendgrid
    from sendgrid.helpers.mail import (Mail, Attachment, FileContent,
                                        FileName, FileType, Disposition)
    SENDGRID_AVAILABLE = True
except ImportError:
    SENDGRID_AVAILABLE = False

try:
    import nltk
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    for r in ['punkt','stopwords','punkt_tab']:
        nltk.download(r, quiet=True)
    NLP_AVAILABLE = True
except ImportError:
    NLP_AVAILABLE = False

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_PARSE_AVAILABLE = True
except ImportError:
    DOCX_PARSE_AVAILABLE = False


# ── 1. ATS SCORE CHECKER ──────────────────────────────────────────────────────

INDUSTRY_KEYWORDS = {
    "tech":       ["python","javascript","sql","api","cloud","aws","docker","git",
                   "agile","scrum","machine learning","data","backend","frontend",
                   "database","linux","ci/cd","microservices","react","node"],
    "finance":    ["excel","financial analysis","budgeting","forecasting","audit",
                   "compliance","risk management","accounts","reporting","cpa",
                   "quickbooks","sap","ifrs","kpi","variance analysis"],
    "marketing":  ["seo","social media","content","analytics","campaigns","google ads",
                   "facebook ads","email marketing","copywriting","brand","roi",
                   "conversion","crm","hubspot","strategy"],
    "management": ["leadership","team management","strategy","stakeholders",
                   "project management","budgets","performance","kpi",
                   "cross-functional","pmo","change management","p&l"],
    "general":    ["communication","problem solving","teamwork","results","analysis",
                   "management","coordination","presentation","reporting","planning",
                   "customer","service","excel","microsoft office"],
}

def check_ats_score(cv_data, job_description=""):
    cv_parts = [cv_data.get("name",""), cv_data.get("title",""), cv_data.get("summary","")]
    for j in cv_data.get("experience",[]): cv_parts += [j.get("job_title",""), j.get("responsibilities","")]
    for e in cv_data.get("education",[]): cv_parts += [e.get("qualification",""), e.get("field","")]
    cv_parts += cv_data.get("skills",[])
    cv_text = " ".join(cv_parts).lower()

    score=0; issues=[]; suggestions=[]; found_kw=[]; missing_kw=[]

    fmt=0
    if cv_data.get("email"):    fmt+=5
    else: issues.append("Missing email address")
    if cv_data.get("phone"):    fmt+=5
    else: issues.append("Missing phone number")
    if cv_data.get("location"): fmt+=3
    else: issues.append("Missing city/location")
    if cv_data.get("linkedin"):
        fmt+=4; suggestions.append("LinkedIn URL present — great for visibility")
    else: suggestions.append("Add your LinkedIn URL to boost ATS ranking")
    summary=cv_data.get("summary","")
    if len(summary)>=100:  fmt+=5
    elif summary:          fmt+=2; suggestions.append("Expand your summary to 100+ characters")
    else: issues.append("No professional summary — ATS systems expect one")
    exp=cv_data.get("experience",[])
    fmt += 5 if len(exp)>=2 else (3 if exp else 0)
    if cv_data.get("education"): fmt+=3
    score+=fmt

    skills=cv_data.get("skills",[]); sk=0
    if   len(skills)>=8: sk+=15
    elif len(skills)>=5: sk+=10; suggestions.append(f"Add more skills — you have {len(skills)}, aim for 8-12")
    elif skills:         sk+=5;  suggestions.append("Add more skills — aim for 8-12")
    else: issues.append("No skills listed — ATS filters heavily on skills")

    tl=cv_data.get("title","").lower()
    industry="general"
    if any(w in tl for w in ["engineer","developer","data","software","tech","it"]): industry="tech"
    elif any(w in tl for w in ["finance","accountant","analyst","audit"]): industry="finance"
    elif any(w in tl for w in ["market","brand","content","social"]): industry="marketing"
    elif any(w in tl for w in ["manager","director","head","lead"]): industry="management"

    for kw in INDUSTRY_KEYWORDS.get(industry, INDUSTRY_KEYWORDS["general"]):
        if kw in cv_text: found_kw.append(kw); sk+=1
        else: missing_kw.append(kw)
    score+=min(sk,30)

    jd_score=20
    if job_description:
        if NLP_AVAILABLE:
            try:
                vec=TfidfVectorizer(stop_words="english",max_features=100)
                mat=vec.fit_transform([cv_text, job_description.lower()])
                sim=cosine_similarity(mat[0:1],mat[1:2])[0][0]
                jd_score=int(sim*40)
                stop={'with','that','this','have','will','your','from','they','been',
                      'more','also','into','than','then','when','some','must','should'}
                gap=list(set(re.findall(r'\b[a-z]{4,}\b',job_description.lower()))-
                          set(re.findall(r'\b[a-z]{4,}\b',cv_text))-stop)[:8]
                if gap: suggestions.append(f"Add these keywords from the JD: {', '.join(gap)}")
            except: jd_score=20
        else:
            jd_words=set(job_description.lower().split()); cv_words=set(cv_text.split())
            jd_score=int(len(jd_words&cv_words)/max(len(jd_words),1)*40)
    else:
        suggestions.append("Paste a job description for a precise match score")
    score+=jd_score; score=min(score,100)

    if   score>=80: grade,label,color="A","Excellent — Strong ATS pass rate","#0F6E56"
    elif score>=65: grade,label,color="B","Good — Minor improvements needed","#2196F3"
    elif score>=50: grade,label,color="C","Average — Several gaps to fix","#FF9800"
    else:           grade,label,color="D","Weak — Needs significant work","#f44336"

    return {"score":score,"grade":grade,"grade_label":label,"grade_color":color,
            "formatting_score":fmt,"skills_score":min(sk,30),"jd_score":jd_score,
            "issues":issues,"suggestions":suggestions,"found_keywords":found_kw[:10],
            "missing_keywords":missing_kw[:6],"industry_detected":industry}


# ── 2. AI CV ENHANCER ─────────────────────────────────────────────────────────

def enhance_summary(cv_data):
    if not OPENAI_AVAILABLE: raise RuntimeError("pip install openai")
    key=os.getenv("OPENAI_API_KEY")
    if not key: raise RuntimeError("OPENAI_API_KEY not set in .env")
    client=OpenAI(api_key=key)
    exp_text="\n".join(f"- {j.get('job_title','?')} at {j.get('company','?')}: {j.get('responsibilities','')[:150]}"
                        for j in cv_data.get("experience",[]))
    prompt=f"""Improve this CV. Return ONLY valid JSON, no markdown.
Name:{cv_data.get('name','')} Title:{cv_data.get('title','')}
Summary:{cv_data.get('summary','None')}
Experience:{exp_text}
Skills:{', '.join(cv_data.get('skills',[]))}
JSON format:
{{"summary":"improved 2-3 sentence summary max 280 chars",
  "experience":[{{"responsibilities":"3 bullet points with action verbs separated by newlines"}}]}}"""
    resp=client.chat.completions.create(model="gpt-3.5-turbo",
         messages=[{"role":"user","content":prompt}],temperature=0.7,max_tokens=600)
    raw=resp.choices[0].message.content.strip().replace("```json","").replace("```","").strip()
    try: improved=json.loads(raw)
    except: m=re.search(r'"summary"\s*:\s*"([^"]+)"',raw); improved={"summary":m.group(1) if m else cv_data.get("summary","")}
    enhanced=dict(cv_data)
    if "summary" in improved: enhanced["summary"]=improved["summary"][:280]
    if "experience" in improved and isinstance(improved["experience"],list):
        for i,imp in enumerate(improved["experience"]):
            if i<len(enhanced["experience"]) and "responsibilities" in imp:
                enhanced["experience"][i]=dict(enhanced["experience"][i])
                enhanced["experience"][i]["responsibilities"]=imp["responsibilities"]
    return enhanced


# ── 3. COVER LETTER GENERATOR ─────────────────────────────────────────────────

def generate_cover_letter(cv_data, job_title="", company=""):
    if not OPENAI_AVAILABLE: raise RuntimeError("pip install openai")
    key=os.getenv("OPENAI_API_KEY")
    if not key: raise RuntimeError("OPENAI_API_KEY not set in .env")
    client=OpenAI(api_key=key)
    exp=", ".join(f"{j.get('job_title','')} at {j.get('company','')}" for j in cv_data.get("experience",[])[:2])
    prompt=f"""Write a professional cover letter.
Applicant:{cv_data.get('name','')} Title:{cv_data.get('title','')}
Applying for:{'the '+job_title+' role' if job_title else 'this position'} {'at '+company if company else ''}
Experience:{exp} Skills:{', '.join(cv_data.get('skills',[])[:8])}
Summary:{cv_data.get('summary','')}
Rules: 3 paragraphs (hook/proof/CTA), professional, under 300 words, no placeholders.
Start: Dear Hiring Manager,  End: Sincerely,\\n{cv_data.get('name','')}"""
    resp=client.chat.completions.create(model="gpt-3.5-turbo",
         messages=[{"role":"user","content":prompt}],temperature=0.8,max_tokens=500)
    return resp.choices[0].message.content.strip()


# ── 4. QR CODE GENERATOR ──────────────────────────────────────────────────────

def generate_qr_b64(url, size=150):
    if not QR_AVAILABLE: raise RuntimeError("pip install qrcode[pil]")
    qr=qrcode.QRCode(version=1,error_correction=qrcode.constants.ERROR_CORRECT_M,box_size=4,border=2)
    qr.add_data(url); qr.make(fit=True)
    img=qr.make_image(fill_color="black",back_color="white").resize((size,size))
    buf=io.BytesIO(); img.save(buf,format="PNG")
    return base64.b64encode(buf.getvalue()).decode()


# ── 5. EMAIL CV ───────────────────────────────────────────────────────────────

def email_cv(to_email, pdf_bytes, person_name):
    if not SENDGRID_AVAILABLE: raise RuntimeError("pip install sendgrid")
    key=os.getenv("SENDGRID_API_KEY")
    if not key: raise RuntimeError("SENDGRID_API_KEY not set in .env")
    from_email=os.getenv("SENDGRID_FROM_EMAIL","cv@converterpro.com")
    html=f"""<div style="font-family:Arial,sans-serif;max-width:600px;margin:0 auto">
<div style="background:#0F6E56;padding:28px;text-align:center"><h1 style="color:white;margin:0">ConverterPro</h1></div>
<div style="padding:32px"><h2 style="color:#0F6E56">Your CV is ready, {person_name}!</h2>
<p style="color:#555;line-height:1.7">Your professional CV is attached. Good luck!</p>
<ul style="color:#555;line-height:2"><li>Tailor keywords to each job description</li>
<li>Keep your LinkedIn profile updated</li><li>Follow up within 5-7 days</li></ul>
<div style="text-align:center;margin-top:28px">
<a href="https://converterpro.com" style="background:#0F6E56;color:white;padding:12px 28px;
text-decoration:none;border-radius:6px;font-weight:700">Build Another CV</a></div></div>
<div style="padding:16px;text-align:center;color:#999;font-size:12px">ConverterPro</div></div>"""
    msg=Mail(from_email=from_email,to_emails=to_email,
             subject=f"Your ConverterPro CV — {person_name}",html_content=html)
    att=Attachment(FileContent(base64.b64encode(pdf_bytes).decode()),
                   FileName(f"{person_name.replace(' ','_')}_CV.pdf"),
                   FileType("application/pdf"),Disposition("attachment"))
    msg.attachment=att
    r=sendgrid.SendGridAPIClient(api_key=key).send(msg)
    return r.status_code in (200,202)


# ── 6. RESUME PARSER ──────────────────────────────────────────────────────────

def parse_resume(file_path):
    ext=os.path.splitext(file_path)[1].lower(); text=""
    if ext==".pdf":
        if not PYPDF_AVAILABLE: raise RuntimeError("pypdf not installed")
        reader=pypdf.PdfReader(file_path)
        text="\n".join(p.extract_text() or "" for p in reader.pages)
    elif ext==".docx":
        if not DOCX_PARSE_AVAILABLE: raise RuntimeError("python-docx not installed")
        doc=DocxDocument(file_path)
        text="\n".join(p.text for p in doc.paragraphs)
    elif ext==".txt":
        with open(file_path,encoding="utf-8",errors="replace") as f: text=f.read()
    else: raise ValueError(f"Unsupported format: {ext}")
    return _extract_cv_fields(text)

def _extract_cv_fields(text):
    lines=[l.strip() for l in text.splitlines() if l.strip()]
    result={}
    m=re.search(r'[\w.\-+]+@[\w.\-]+\.\w+',text)
    if m: result["email"]=m.group(0)
    m=re.search(r'(\+?254|0)[\s\-]?[7][0-9][\s\-]?\d{3}[\s\-]?\d{4}',text)
    if m: result["phone"]=re.sub(r'[\s\-]','',m.group(0))
    m=re.search(r'linkedin\.com/in/[\w\-]+',text,re.IGNORECASE)
    if m: result["linkedin"]="https://"+m.group(0)
    m=re.search(r'https?://(?!linkedin)[\w.\-/]+',text,re.IGNORECASE)
    if m: result["website"]=m.group(0)
    for line in lines[:5]:
        if re.match(r'^[A-Z][a-z]+ [A-Z][a-z]+',line) and len(line)<60:
            result["name"]=line; break
    name_idx=next((i for i,l in enumerate(lines[:8]) if result.get("name","") in l),0)
    for line in lines[name_idx+1:name_idx+5]:
        if not re.search(r'[@|/]|\d{6,}',line) and 4<len(line)<80:
            result["title"]=line; break
    skills=[]; in_skills=False
    for line in lines:
        if re.match(r'^skills?',line,re.IGNORECASE): in_skills=True; continue
        if in_skills:
            if re.match(r'^(experience|education|summary|work|employment|projects)',line,re.IGNORECASE): break
            for item in re.split(r'[,|•·\n]',line):
                item=item.strip().strip('●▪-–').strip()
                if 2<len(item)<40: skills.append(item)
    if skills: result["skills_hidden"]=", ".join(skills[:15])
    sum_lines=[]; in_sum=False
    for line in lines:
        if re.match(r'^(summary|profile|objective|about)',line,re.IGNORECASE): in_sum=True; continue
        if in_sum:
            if re.match(r'^(experience|education|skills|work)',line,re.IGNORECASE): break
            if line: sum_lines.append(line)
            if len(" ".join(sum_lines))>250: break
    if sum_lines: result["summary"]=" ".join(sum_lines)[:300]
    return result


def feature_status():
    return {"ai_enhance":OPENAI_AVAILABLE and bool(os.getenv("OPENAI_API_KEY")),
            "cover_letter":OPENAI_AVAILABLE and bool(os.getenv("OPENAI_API_KEY")),
            "ats_score":True,"qr_code":QR_AVAILABLE,
            "email_cv":SENDGRID_AVAILABLE and bool(os.getenv("SENDGRID_API_KEY")),
            "resume_parse":PYPDF_AVAILABLE or DOCX_PARSE_AVAILABLE,
            "packages":{"openai":OPENAI_AVAILABLE,"qrcode":QR_AVAILABLE,
                        "sendgrid":SENDGRID_AVAILABLE,"nltk_sklearn":NLP_AVAILABLE,
                        "pypdf":PYPDF_AVAILABLE,"python_docx":DOCX_PARSE_AVAILABLE}}
